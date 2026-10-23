"""Dissect an .xlsx with a real reader and print a JSON report.

Used by test_xlsx_writer.js so "the workbook is valid" is checked by openpyxl
plus the raw OOXML parts, not by our own writer agreeing with itself.
"""
import json
import re
import struct
import sys
import zipfile


def jpeg_size(b):
    i = 2
    while i < len(b) - 9:
        if b[i] != 0xFF:
            i += 1
            continue
        m = b[i + 1]
        if m in (0xC0, 0xC1, 0xC2, 0xC3):
            h, w = struct.unpack('>HH', b[i + 5:i + 9])
            return [w, h]
        if m in (0xD8, 0xD9) or 0xD0 <= m <= 0xD7:
            i += 2
            continue
        i += 2 + struct.unpack('>H', b[i + 2:i + 4])[0]
    return None


def main(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    rep = {'sheets': wb.sheetnames}

    summary = wb[wb.sheetnames[0]]
    pairs = {}
    for row in summary.iter_rows(values_only=True):
        if row and row[0] and len(row) > 1 and row[1] is not None:
            pairs[str(row[0])] = row[1]
    rep['summary_pairs'] = pairs

    ws = wb[wb.sheetnames[-1]]
    headers = [c.value for c in ws[1] if c.value is not None]
    rep['headers'] = headers
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and v != '' for v in r):
            continue
        rows.append({h: ('' if v is None else v) for h, v in zip(headers, r)})
    rep['rows'] = rows
    rep['row_count'] = len(rows)
    rep['row_heights'] = [ws.row_dimensions[i].height for i in range(2, 2 + len(rows))]
    rep['col_widths'] = {k: v.width for k, v in ws.column_dimensions.items() if v.width}

    # Raw parts: drawings, media and hyperlink rels.
    z = zipfile.ZipFile(path)
    names = z.namelist()
    media = sorted([n for n in names if n.startswith('xl/media/') and not n.endswith('/')])
    rep['image_count'] = len(media)
    rep['image_sizes'] = [jpeg_size(z.read(n)) for n in media]

    last = len(wb.sheetnames)
    drawing = 'xl/drawings/drawing%d.xml' % last
    cols, rws = [], []
    if drawing in names:
        d = z.read(drawing).decode('utf8')
        cols = [int(x) for x in re.findall(r'<xdr:col>(\d+)</xdr:col>', d)]
        rws = [int(x) for x in re.findall(r'<xdr:row>(\d+)</xdr:row>', d)]
    rep['anchor_cols'] = cols
    rep['anchor_rows'] = rws
    rep['summary_images'] = 1 if 'xl/drawings/drawing1.xml' in names and last != 1 else 0

    links = {}
    sheet_xml = 'xl/worksheets/sheet%d.xml' % last
    rels_name = 'xl/worksheets/_rels/sheet%d.xml.rels' % last
    if sheet_xml in names and rels_name in names:
        sx = z.read(sheet_xml).decode('utf8')
        rx = z.read(rels_name).decode('utf8')
        targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rx))
        for ref, rid in re.findall(r'<hyperlink ref="([^"]+)" r:id="([^"]+)"/>', sx):
            t = targets.get(rid, '')
            links[ref] = (t.replace('&amp;', '&').replace('&lt;', '<')
                           .replace('&gt;', '>').replace('&quot;', '"'))
    rep['hyperlinks'] = links
    rep['hyperlink_count'] = len(links)
    return rep


if __name__ == '__main__':
    try:
        print(json.dumps(main(sys.argv[1])))
    except Exception as e:  # surfaced as a test failure, not a crash
        print(json.dumps({'error': '%s: %s' % (type(e).__name__, e)}))
